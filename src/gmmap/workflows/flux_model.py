#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
metabolic_flow_model_wAUCellS3_ratio_enzyme_constrained_round13.py

Hybrid metabolic-flow model:
    metabolite-linked wAUCell-S3 scores
    + ratio-based direction support
    + enzyme-capacity soft constraint
    + compartment-support soft constraint
    + cofactor-support soft constraint

This is NOT a full FBA implementation. Instead, it is a transcriptome-driven
flux-potential model with Compass-inspired feasibility constraints.

Inputs
------
1) annotation table (round12 csv/xlsx)
   required columns (auto-detected):
   - module_id
   - trait_in
   - trait_out
   optional:
   - ratio_trait
   - direction_sign
   - subpathway_en / subpathway_cn
   - pathway_class
   - in_name / out_name

2) wAUCell-S3 metabolite score matrix
   rows = cells or celltypes
   cols = metabolite traits (GCST...)

3) gene expression matrix
   rows = cells or celltypes
   cols = genes

4) optional metadata
   used to aggregate both matrices to celltype

5) optional enzyme annotation
   can be:
   - round7 xlsx workbook
   - csv/xlsx with module_id + enzyme genes columns
   If not provided, only built-in compartment/cofactor heuristics are used.

Model
-----
For module m in celltype c:

s_in, s_out, s_ratio = metabolite / ratio scores

activation ~ tanh(mean(|s_in|, |s_out|))
direction  ~ tanh((s_out - s_in) + ratio_weight * direction_sign * s_ratio)

enzyme_capacity       = GPR-like proxy from enzyme genes
compartment_support   = soft support from compartment hallmark genes
cofactor_support      = soft support from pathway/cofactor hallmark genes

feasibility =
    node_support *
    coverage *
    ratio_factor *
    consistency *
    enzyme_capacity^enzyme_weight *
    compartment_support^compartment_weight *
    cofactor_support^cofactor_weight

signed_flux = activation * direction * confidence
where confidence = feasibility

Outputs
-------
Module-level:
- module_activation_matrix.csv.gz
- module_direction_matrix.csv.gz
- module_confidence_matrix.csv.gz
- module_signed_flux_matrix.csv.gz
- module_ratio_support_matrix.csv.gz
- module_enzyme_capacity_matrix.csv.gz
- module_compartment_support_matrix.csv.gz
- module_cofactor_support_matrix.csv.gz
- module_feasibility_matrix.csv.gz
- module_mean_s3_matrix.csv.gz
- module_delta_s3_matrix.csv.gz

Subpathway-level:
- subpathway_activation_matrix.csv.gz
- subpathway_direction_matrix.csv.gz
- subpathway_confidence_matrix.csv.gz
- subpathway_signed_flux_matrix.csv.gz
- subpathway_ratio_support_matrix.csv.gz
- subpathway_enzyme_capacity_matrix.csv.gz
- subpathway_compartment_support_matrix.csv.gz
- subpathway_cofactor_support_matrix.csv.gz
- subpathway_feasibility_matrix.csv.gz
- subpathway_mean_s3_matrix.csv.gz
- subpathway_delta_s3_matrix.csv.gz

Diagnostics:
- module_scores_long.csv.gz
- subpathway_scores_long.csv.gz
- module_annotation_used.csv
- module_missing_trait_report.csv
- module_constraint_gene_report.csv
- run_summary.csv
"""

import argparse
import re
from pathlib import Path
from typing import Dict, List, Tuple, Optional

import numpy as np
import pandas as pd

try:
    import openpyxl
except Exception:
    openpyxl = None


# ----------------------------
# Utilities
# ----------------------------

def tanh(x: float) -> float:
    return float(np.tanh(x))


def mean_ignore_nan(vals):
    vals = [float(v) for v in vals if pd.notna(v) and np.isfinite(v)]
    if len(vals) == 0:
        return np.nan
    return float(np.mean(vals))


def weighted_mean(values, weights):
    v = np.asarray(values, dtype=float)
    w = np.asarray(weights, dtype=float)
    m = np.isfinite(v) & np.isfinite(w) & (w > 0)
    if m.sum() == 0:
        return np.nan
    return float(np.sum(v[m] * w[m]) / np.sum(w[m]))


def safe_str(x):
    return "" if pd.isna(x) else str(x).strip()


def normalize_gene_symbol(g: str) -> str:
    return re.sub(r"\s+", "", str(g).strip()).upper()


def topk_mean(vals: List[float], k: int = 3) -> float:
    vals = [float(v) for v in vals if pd.notna(v) and np.isfinite(v)]
    if len(vals) == 0:
        return np.nan
    vals = sorted(vals, reverse=True)
    return float(np.mean(vals[:min(k, len(vals))]))


def geometric_mean_soft(vals: List[float]) -> float:
    vals = [float(v) for v in vals if pd.notna(v) and np.isfinite(v) and v >= 0]
    if len(vals) == 0:
        return np.nan
    vals = np.asarray(vals, dtype=float)
    return float(np.exp(np.mean(np.log(vals + 1e-8))))


# ----------------------------
# Reading inputs
# ----------------------------

def read_table(path: str) -> pd.DataFrame:
    p = Path(path)
    if p.suffix.lower() in [".xlsx", ".xls", ".xlsm"]:
        return pd.read_excel(p)
    return pd.read_csv(p, compression="infer")


def read_matrix(path: str) -> pd.DataFrame:
    p = Path(path)
    if p.suffix.lower() in [".xlsx", ".xls", ".xlsm"]:
        df = pd.read_excel(p, index_col=0)
    else:
        df = pd.read_csv(p, index_col=0, compression="infer")
    df.index = df.index.astype(str)
    df.columns = df.columns.astype(str)
    return df


def aggregate_to_celltype(df_mat, metadata_path=None, id_col=None, celltype_col=None, agg="mean"):
    if metadata_path is None:
        out = df_mat.copy()
        out.index = out.index.astype(str)
        return out

    meta = pd.read_csv(metadata_path)
    if id_col is None or celltype_col is None:
        raise ValueError("When --metadata is provided, --id-col and --celltype-col are required.")
    if id_col not in meta.columns or celltype_col not in meta.columns:
        raise ValueError(f"metadata missing columns: {id_col}, {celltype_col}")

    meta = meta[[id_col, celltype_col]].copy()
    meta[id_col] = meta[id_col].astype(str)
    meta[celltype_col] = meta[celltype_col].astype(str)
    common = meta[id_col].isin(df_mat.index)
    meta = meta.loc[common].copy()
    if meta.empty:
        raise ValueError("No metadata row IDs overlap matrix index.")
    df = df_mat.loc[meta[id_col]].copy()
    df["__celltype__"] = meta.set_index(id_col).loc[df.index, celltype_col].values

    if agg == "median":
        out = df.groupby("__celltype__").median(numeric_only=True)
    else:
        out = df.groupby("__celltype__").mean(numeric_only=True)
    out.index = out.index.astype(str)
    return out


# ----------------------------
# Annotation
# ----------------------------

def detect_annotation_columns(df):
    cols = list(df.columns)
    lookup = {str(c).lower(): c for c in cols}

    def first_match(candidates):
        for cand in candidates:
            if cand in cols:
                return cand
            if cand.lower() in lookup:
                return lookup[cand.lower()]
        return None

    out = {
        "module_id": first_match(["module_id", "Module_id", "列1"]),
        "trait_in": first_match(["trait_in", "accessionId_in", "accessionid_in"]),
        "trait_out": first_match(["trait_out", "accessionId_out", "accessionid_out"]),
        "ratio_trait": first_match(["ratio_trait"]),
        "direction_sign": first_match(["direction_sign"]),
        "subpathway_cn": first_match(["subpathway_cn", "Subpathway", "subpathway"]),
        "subpathway_en": first_match(["subpathway_en"]),
        "pathway_class": first_match(["pathway_class", "Pathway_class"]),
        "in_name": first_match(["in_name", "Compound_IN_name", "compound_in_name"]),
        "out_name": first_match(["out_name", "Compound_OUT_name", "compound_out_name"]),
        "rule_confidence": first_match(["rule_confidence", "Rule_confidence"]),
        "note": first_match(["note"]),
    }
    if any(out[k] is None for k in ["module_id", "trait_in", "trait_out"]):
        raise ValueError(
            f"Cannot auto-detect required annotation columns. "
            f"Detected module_id={out['module_id']}, trait_in={out['trait_in']}, trait_out={out['trait_out']}"
        )
    return out


def prepare_annotation(df_anno):
    cmap = detect_annotation_columns(df_anno)
    out = pd.DataFrame()
    for key, src in cmap.items():
        out[key] = np.nan if src is None else df_anno[src]
    out["module_id"] = out["module_id"].astype(str).str.strip()
    out["trait_in"] = out["trait_in"].astype(str).str.strip()
    out["trait_out"] = out["trait_out"].astype(str).str.strip()
    out["ratio_trait"] = out["ratio_trait"].replace(["", "None", "nan"], np.nan)
    out.loc[out["ratio_trait"].notna(), "ratio_trait"] = out.loc[out["ratio_trait"].notna(), "ratio_trait"].astype(str).str.strip()
    out["direction_sign"] = pd.to_numeric(out["direction_sign"], errors="coerce")
    out["subpathway_cn"] = out["subpathway_cn"].fillna("").astype(str)
    out["subpathway_en"] = out["subpathway_en"].fillna("").astype(str)
    out.loc[out["subpathway_en"].str.strip() == "", "subpathway_en"] = out.loc[out["subpathway_en"].str.strip() == "", "subpathway_cn"]
    return out


# ----------------------------
# Enzyme annotation parsing
# ----------------------------

def detect_gene_columns(df):
    cols = list(df.columns)
    lookup = {str(c).lower(): c for c in cols}

    def first_match(cands):
        for c in cands:
            if c in cols:
                return c
            if c.lower() in lookup:
                return lookup[c.lower()]
        return None

    return {
        "module_id": first_match(["module_id", "Module_id"]),
        "enzyme_gene_symbols": first_match(["Enzyme_gene_symbols"]),
        "specific_gene_symbols": first_match(["Specific_gene_symbols"]),
        "enzyme_step": first_match(["Enzyme_or_key_step_superdetailed"]),
        "in_name": first_match(["Compound_IN_name", "in_name"]),
        "out_name": first_match(["Compound_OUT_name", "out_name"]),
    }


def load_round7_total_sheet(path: str) -> pd.DataFrame:
    p = Path(path)
    if p.suffix.lower() not in [".xlsx", ".xls", ".xlsm"]:
        return read_table(path)
    if openpyxl is None:
        raise ImportError("openpyxl is required to read xlsx enzyme annotation")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    if "总表_逐代谢物模块反应" in wb.sheetnames:
        ws = wb["总表_逐代谢物模块反应"]
    elif "模块基因信息" in wb.sheetnames:
        ws = wb["模块基因信息"]
    else:
        ws = wb[wb.sheetnames[0]]
    rows = list(ws.values)
    header = [str(x) if x is not None else "" for x in rows[0]]
    return pd.DataFrame(rows[1:], columns=header)


def split_gene_text(s: str) -> List[str]:
    s = safe_str(s)
    if s == "":
        return []
    s = s.replace("；", ";").replace("，", ",").replace("|", ",")
    s = re.sub(r"\band\b", ",", s, flags=re.IGNORECASE)
    # remove descriptors like "gene:" or "候选:"
    s = re.sub(r"[A-Za-z\u4e00-\u9fff\- ]+:\s*", "", s)
    parts = re.split(r"[;,]", s)
    raw = []
    for p in parts:
        p = p.strip()
        if not p:
            continue
        # expand patterns like CPT1A/B/C
        m = re.match(r"^([A-Za-z0-9]+?)([A-Z])(?:/([A-Z]))(?:/([A-Z]))?$", p)
        if m:
            prefix = m.group(1)
            suffs = [x for x in m.groups()[1:] if x]
            for suf in suffs:
                raw.append(prefix + suf)
            continue
        # expand patterns like IDO1/IDO2
        m2 = re.match(r"^([A-Za-z]+)(\d+)(?:/([A-Za-z]+)?(\d+))+$", p)
        if "/" not in p:
            raw.append(p)
            continue
        # fallback slash split
        if "/" in p and len(p) < 40 and "http" not in p:
            # attempt prefix expansion
            tokens = p.split("/")
            if len(tokens) >= 2:
                head = tokens[0]
                if re.match(r"^[A-Za-z]+[0-9]*$", head):
                    prefix = re.match(r"^([A-Za-z]+)", head).group(1)
                    raw.append(head)
                    for tok in tokens[1:]:
                        tok = tok.strip()
                        if re.match(r"^[0-9A-Za-z]+$", tok):
                            if tok[0].isdigit():
                                raw.append(prefix + tok)
                            elif len(tok) <= 3:
                                raw.append(prefix + tok)
                            else:
                                raw.append(tok)
                    continue
        raw.append(p)

    # keep gene-like strings only
    cleaned = []
    for g in raw:
        g = normalize_gene_symbol(g)
        if g == "":
            continue
        # reject obvious non-gene phrases
        if len(g) > 20:
            continue
        if not re.search(r"[A-Z]", g):
            continue
        cleaned.append(g)
    # deduplicate
    out = []
    for g in cleaned:
        if g not in out:
            out.append(g)
    return out


def build_enzyme_map(enzyme_annotation_path: Optional[str]) -> pd.DataFrame:
    if enzyme_annotation_path is None:
        return pd.DataFrame(columns=["module_id", "enzyme_genes", "enzyme_step"])
    df = load_round7_total_sheet(enzyme_annotation_path)
    gc = detect_gene_columns(df)
    if gc["module_id"] is None:
        raise ValueError("Cannot detect module_id in enzyme annotation")
    out = pd.DataFrame()
    out["module_id"] = df[gc["module_id"]].astype(str).str.strip()
    g1 = df[gc["specific_gene_symbols"]] if gc["specific_gene_symbols"] is not None else pd.Series("", index=df.index)
    g2 = df[gc["enzyme_gene_symbols"]] if gc["enzyme_gene_symbols"] is not None else pd.Series("", index=df.index)
    genes = []
    for a, b in zip(g1.fillna(""), g2.fillna("")):
        merged = split_gene_text(str(a)) or split_gene_text(str(b))
        genes.append(merged)
    out["enzyme_genes"] = genes
    out["enzyme_step"] = df[gc["enzyme_step"]] if gc["enzyme_step"] is not None else ""
    if gc["in_name"] is not None:
        out["in_name_ref"] = df[gc["in_name"]].astype(str)
    if gc["out_name"] is not None:
        out["out_name_ref"] = df[gc["out_name"]].astype(str)
    return out.drop_duplicates(subset=["module_id"])


# ----------------------------
# Compass-inspired soft constraints
# ----------------------------

MITO_PANEL = [
    "CPT1A","CPT1B","CPT1C","CPT2","SLC25A20","CRAT","ACADVL","ACADL","ACADM",
    "ECHS1","HADHA","HADHB","ACAA2","CS","ACO2","IDH3A","OGDH","SUCLG1","SDHA","FH","MDH2","MPC1","MPC2"
]
PEROX_PANEL = [
    "ABCD1","ABCD2","ABCD3","ACOX1","HSD17B4","CROT","ACAA1","SCP2","PEX5","PEX13"
]
CYTOSOL_GLYCOLYSIS_PANEL = [
    "HK1","HK2","GPI","PFKL","PFKM","PFKP","ALDOA","GAPDH","PGK1","ENO1","PKM","LDHA","LDHB","G6PD","PGD"
]
UREA_PANEL = ["CPS1","OTC","ASS1","ASL","ARG1","ARG2","SLC25A15"]
TRP_KYN_PANEL = ["IDO1","IDO2","TDO2","AFMID","KMO","KYNU","KYAT1","AADAT","CCBL2","GOT2"]
TRANSPORT_PANEL = ["SLC22A5","SLC25A20","ABCD1","ABCD2","ABCD3","MPC1","MPC2","SLC25A15"]

FAO_COF_PANEL = ["ETFA","ETFB","ETFDH","NDUFS1","NDUFA9","UQCRC1","COX4I1","ATP5F1A"]
GLYCOLYSIS_COF_PANEL = ["LDHA","LDHB","G6PD","PGD","ME1","ME2","PDHA1","PDHB"]
UREA_COF_PANEL = ["CPS1","ASS1","ASL","ARG1","ARG2"]
TRP_COF_PANEL = ["IDO1","IDO2","TDO2","KMO","AFMID","KYNU"]


def infer_compartment_panels(row) -> List[List[str]]:
    sp = safe_str(row.get("subpathway_en", "")).lower()
    pc = safe_str(row.get("pathway_class", "")).lower()
    txt = " ".join([sp, pc, safe_str(row.get("in_name","")).lower(), safe_str(row.get("out_name","")).lower()])

    panels = []
    if any(k in txt for k in ["acylcarnitine", "carnitine", "tca", "citrate", "aconitate", "malate", "succinate", "α-ketoglutarate", "ketone", "pyruvate entry"]):
        panels.append(MITO_PANEL)
    if any(k in txt for k in ["dicarboxylic", "very-long", "perox", "hydroxyacylcarnitine"]):
        panels.append(PEROX_PANEL)
    if any(k in txt for k in ["lactate", "hexose", "disaccharide", "glycerate", "triose", "glucuronate"]):
        panels.append(CYTOSOL_GLYCOLYSIS_PANEL)
    if any(k in txt for k in ["arginine", "urea", "ornithine", "citrulline"]):
        panels.append(UREA_PANEL)
    if any(k in txt for k in ["tryptophan", "kynurenine", "indole", "histidine"]):
        panels.append(TRP_KYN_PANEL)
    if any(k in txt for k in ["transport", "carnitine", "acylcarnitine"]):
        panels.append(TRANSPORT_PANEL)
    return panels


def infer_cofactor_panel(row) -> List[str]:
    sp = safe_str(row.get("subpathway_en", "")).lower()
    pc = safe_str(row.get("pathway_class", "")).lower()
    txt = " ".join([sp, pc, safe_str(row.get("in_name","")).lower(), safe_str(row.get("out_name","")).lower()])

    if any(k in txt for k in ["acylcarnitine", "carnitine", "fatty acid", "ketone", "tca", "succinate", "malate", "citrate", "α-ketoglutarate"]):
        return FAO_COF_PANEL
    if any(k in txt for k in ["lactate", "pyruvate", "hexose", "glycol", "glycerate", "triose", "glucuronate"]):
        return GLYCOLYSIS_COF_PANEL
    if any(k in txt for k in ["arginine", "urea", "ornithine", "citrulline"]):
        return UREA_COF_PANEL
    if any(k in txt for k in ["tryptophan", "kynurenine", "indole", "histidine"]):
        return TRP_COF_PANEL
    return []


def score_gene_set(expr_row: pd.Series, genes: List[str], mode: str = "topk_mean", topk: int = 3, scale: float = 1.0):
    genes = [normalize_gene_symbol(g) for g in genes]
    avail = [g for g in genes if g in expr_row.index]
    vals = [expr_row[g] for g in avail if pd.notna(expr_row[g]) and np.isfinite(expr_row[g])]
    if len(avail) == 0:
        return np.nan, 0, len(genes), []

    if len(vals) == 0:
        raw = np.nan
    elif mode == "max":
        raw = float(np.max(vals))
    elif mode == "mean":
        raw = float(np.mean(vals))
    elif mode == "min":
        raw = float(np.min(vals))
    elif mode == "geom":
        raw = geometric_mean_soft(vals)
    else:
        raw = topk_mean(vals, k=topk)

    score = np.nan if not np.isfinite(raw) else float(np.tanh(raw / scale))
    return score, len(avail), len(genes), avail


def combine_panel_scores(scores: List[float]) -> float:
    scores = [float(s) for s in scores if pd.notna(s) and np.isfinite(s)]
    if len(scores) == 0:
        return np.nan
    return float(np.mean(scores))


def compute_consistency(base_direction, ratio_support, zero_eps=1e-8):
    if not np.isfinite(ratio_support):
        return 1.0
    if (not np.isfinite(base_direction)) or abs(base_direction) < zero_eps or abs(ratio_support) < zero_eps:
        return 0.5
    return 1.0 if np.sign(base_direction) == np.sign(ratio_support) else 0.25


# ----------------------------
# Main scoring
# ----------------------------

def merge_annotation_with_enzymes(anno: pd.DataFrame, enz: pd.DataFrame) -> pd.DataFrame:
    if enz.empty:
        anno["enzyme_genes"] = [[] for _ in range(len(anno))]
        anno["enzyme_step_ref"] = np.nan
        return anno
    out = anno.merge(enz[["module_id", "enzyme_genes", "enzyme_step"]], on="module_id", how="left")
    out["enzyme_genes"] = out["enzyme_genes"].apply(lambda x: x if isinstance(x, list) else [])
    out["enzyme_step_ref"] = out["enzyme_step"]
    out = out.drop(columns=["enzyme_step"])
    return out


def compute_module_scores(
    anno: pd.DataFrame,
    score_mat: pd.DataFrame,
    gene_mat: pd.DataFrame,
    ratio_weight: float = 1.0,
    activation_scale: float = 1.0,
    direction_scale: float = 1.0,
    signal_scale: float = 1.0,
    ratio_scale: float = 1.0,
    enzyme_scale: float = 1.0,
    compartment_scale: float = 1.0,
    cofactor_scale: float = 1.0,
    ratio_missing_penalty: float = 0.5,
    no_ratio_factor: float = 1.0,
    no_enzyme_factor: float = 0.8,
    no_compartment_factor: float = 0.9,
    no_cofactor_factor: float = 0.9,
    enzyme_weight: float = 1.0,
    compartment_weight: float = 0.5,
    cofactor_weight: float = 0.25,
    gpr_mode: str = "topk_mean",
    gpr_topk: int = 3,
):
    celltypes = list(score_mat.index)
    if set(celltypes) != set(gene_mat.index):
        common = score_mat.index.intersection(gene_mat.index)
        score_mat = score_mat.loc[common].copy()
        gene_mat = gene_mat.loc[common].copy()
        celltypes = list(common)

    traits_available = set(score_mat.columns)
    gene_mat.columns = pd.Index([normalize_gene_symbol(c) for c in gene_mat.columns])

    rows = []
    missing_rows = []
    constraint_rows = []

    for _, row in anno.iterrows():
        module_id = row["module_id"]
        tr_in = row["trait_in"]
        tr_out = row["trait_out"]
        ratio_trait = row.get("ratio_trait", np.nan)
        direction_sign = row.get("direction_sign", np.nan)
        enz_genes = row.get("enzyme_genes", [])
        enz_genes = enz_genes if isinstance(enz_genes, list) else []

        in_present = tr_in in traits_available
        out_present = tr_out in traits_available
        ratio_assigned = pd.notna(ratio_trait) and str(ratio_trait).strip() != ""
        ratio_present = ratio_assigned and (str(ratio_trait).strip() in traits_available)
        coverage = ((1 if in_present else 0) + (1 if out_present else 0)) / 2.0

        if (not in_present) or (not out_present) or (ratio_assigned and not ratio_present):
            missing_rows.append({
                "module_id": module_id,
                "trait_in": tr_in,
                "trait_out": tr_out,
                "ratio_trait": ratio_trait,
                "in_present": in_present,
                "out_present": out_present,
                "ratio_assigned": ratio_assigned,
                "ratio_present": ratio_present,
                "subpathway_en": row.get("subpathway_en", ""),
            })

        comp_panels = infer_compartment_panels(row)
        cof_panel = infer_cofactor_panel(row)

        for ct in celltypes:
            s_in = score_mat.at[ct, tr_in] if in_present else np.nan
            s_out = score_mat.at[ct, tr_out] if out_present else np.nan

            activation_raw = mean_ignore_nan([abs(s_in), abs(s_out)])
            activation = np.nan if not np.isfinite(activation_raw) else tanh(activation_raw / activation_scale)

            mean_s3 = mean_ignore_nan([s_in, s_out])
            delta_s3 = np.nan if (not np.isfinite(s_in) or not np.isfinite(s_out)) else float(s_out - s_in)
            base_direction = delta_s3

            s_ratio = np.nan
            ratio_support = np.nan
            if ratio_present:
                s_ratio = score_mat.at[ct, str(ratio_trait).strip()]
                if np.isfinite(s_ratio):
                    ds = 1.0 if not np.isfinite(direction_sign) else float(direction_sign)
                    ratio_support = ds * float(s_ratio)

            if np.isfinite(base_direction) and np.isfinite(ratio_support):
                direction_input = base_direction + ratio_weight * ratio_support
            else:
                direction_input = base_direction
            direction = np.nan if not np.isfinite(direction_input) else tanh(direction_input / direction_scale)

            node_support = mean_ignore_nan([
                np.tanh(abs(float(s_in)) / signal_scale) if np.isfinite(s_in) else np.nan,
                np.tanh(abs(float(s_out)) / signal_scale) if np.isfinite(s_out) else np.nan,
            ])

            if ratio_assigned:
                if ratio_present and np.isfinite(s_ratio):
                    ratio_factor = float(np.tanh(abs(float(s_ratio)) / ratio_scale))
                else:
                    ratio_factor = float(ratio_missing_penalty)
            else:
                ratio_factor = float(no_ratio_factor)

            consistency = compute_consistency(base_direction, ratio_support)

            # enzyme capacity
            expr_row = gene_mat.loc[ct]
            enz_score, enz_avail, enz_total, enz_used = score_gene_set(
                expr_row, enz_genes, mode=gpr_mode, topk=gpr_topk, scale=enzyme_scale
            )
            enzyme_coverage = np.nan if enz_total == 0 else float(enz_avail / max(enz_total, 1))
            if enz_total == 0:
                enzyme_capacity = float(no_enzyme_factor)
            else:
                enzyme_capacity = float(no_enzyme_factor if not np.isfinite(enz_score) else (0.5 * enz_score + 0.5 * enzyme_coverage))

            # compartment support
            panel_scores = []
            comp_used = []
            for panel in comp_panels:
                sc, avail, total, used = score_gene_set(
                    expr_row, panel, mode="topk_mean", topk=min(3, len(panel)), scale=compartment_scale
                )
                if np.isfinite(sc):
                    panel_scores.append(sc)
                comp_used.extend(used)
            if len(comp_panels) == 0:
                compartment_support = float(no_compartment_factor)
            else:
                panel_mean = combine_panel_scores(panel_scores)
                compartment_support = float(no_compartment_factor if not np.isfinite(panel_mean) else panel_mean)

            # cofactor support
            cof_score, cof_avail, cof_total, cof_used = score_gene_set(
                expr_row, cof_panel, mode="topk_mean", topk=min(3, len(cof_panel)), scale=cofactor_scale
            )
            if len(cof_panel) == 0:
                cofactor_support = float(no_cofactor_factor)
            else:
                cofactor_support = float(no_cofactor_factor if not np.isfinite(cof_score) else cof_score)

            # feasibility/confidence
            if not np.isfinite(node_support):
                confidence = np.nan
                feasibility = np.nan
            else:
                feasibility = float(
                    coverage
                    * node_support
                    * ratio_factor
                    * consistency
                    * (max(enzyme_capacity, 1e-6) ** enzyme_weight)
                    * (max(compartment_support, 1e-6) ** compartment_weight)
                    * (max(cofactor_support, 1e-6) ** cofactor_weight)
                )
                confidence = feasibility

            signed_flux = (
                float(activation * direction * confidence)
                if np.isfinite(activation) and np.isfinite(direction) and np.isfinite(confidence)
                else np.nan
            )

            rows.append({
                "module_id": module_id,
                "celltype": ct,
                "trait_in": tr_in,
                "trait_out": tr_out,
                "ratio_trait": ratio_trait,
                "direction_sign": direction_sign,
                "subpathway_cn": row.get("subpathway_cn", ""),
                "subpathway_en": row.get("subpathway_en", ""),
                "pathway_class": row.get("pathway_class", np.nan),
                "in_name": row.get("in_name", np.nan),
                "out_name": row.get("out_name", np.nan),
                "rule_confidence": row.get("rule_confidence", np.nan),
                "note": row.get("note", np.nan),
                "enzyme_step_ref": row.get("enzyme_step_ref", np.nan),
                "enzyme_genes": ",".join(enz_genes),
                "signal_in": s_in,
                "signal_out": s_out,
                "signal_ratio": s_ratio,
                "mean_s3": mean_s3,
                "delta_s3": delta_s3,
                "activation": activation,
                "base_direction": base_direction,
                "ratio_support": ratio_support,
                "direction_input": direction_input,
                "direction": direction,
                "coverage": coverage,
                "node_support": node_support,
                "ratio_factor": ratio_factor,
                "consistency": consistency,
                "enzyme_capacity": enzyme_capacity,
                "enzyme_coverage": enzyme_coverage,
                "compartment_support": compartment_support,
                "cofactor_support": cofactor_support,
                "feasibility": feasibility,
                "confidence": confidence,
                "signed_flux": signed_flux,
                "ratio_assigned": ratio_assigned,
                "ratio_present": ratio_present,
            })

            constraint_rows.append({
                "module_id": module_id,
                "celltype": ct,
                "enzyme_genes_total": enz_total,
                "enzyme_genes_available": enz_avail,
                "enzyme_genes_used": ",".join(sorted(set(enz_used))),
                "compartment_genes_used": ",".join(sorted(set(comp_used))),
                "cofactor_genes_used": ",".join(sorted(set(cof_used))),
                "enzyme_capacity": enzyme_capacity,
                "compartment_support": compartment_support,
                "cofactor_support": cofactor_support,
            })

    module_long = pd.DataFrame(rows)
    missing_df = pd.DataFrame(missing_rows).drop_duplicates()
    constraint_df = pd.DataFrame(constraint_rows)
    return module_long, missing_df, constraint_df


def pivot_metric(df_long, value_col, index_col="module_id", column_col="celltype"):
    return df_long.pivot(index=index_col, columns=column_col, values=value_col)


def aggregate_subpathway(module_long):
    rows = []
    for (sp_cn, sp_en, ct), g in module_long.groupby(["subpathway_cn", "subpathway_en", "celltype"], dropna=False):
        weights = g["confidence"].to_numpy(dtype=float)
        sp_cn = "" if pd.isna(sp_cn) else str(sp_cn)
        sp_en = "" if pd.isna(sp_en) else str(sp_en)
        if sp_en.strip() and sp_cn.strip() and sp_en.strip() != sp_cn.strip():
            subpathway_key = f"{sp_en} | {sp_cn}"
        else:
            subpathway_key = sp_en.strip() if sp_en.strip() else sp_cn.strip()
        rows.append({
            "subpathway_key": subpathway_key,
            "subpathway_cn": sp_cn,
            "subpathway_en": sp_en,
            "celltype": ct,
            "n_modules": int(g["module_id"].nunique()),
            "activation": weighted_mean(g["activation"], weights),
            "direction": weighted_mean(g["direction"], weights),
            "confidence": mean_ignore_nan(g["confidence"].tolist()),
            "signed_flux": weighted_mean(g["signed_flux"], weights),
            "ratio_support": weighted_mean(g["ratio_support"], weights),
            "mean_s3": weighted_mean(g["mean_s3"], weights),
            "delta_s3": weighted_mean(g["delta_s3"], weights),
            "enzyme_capacity": weighted_mean(g["enzyme_capacity"], weights),
            "compartment_support": weighted_mean(g["compartment_support"], weights),
            "cofactor_support": weighted_mean(g["cofactor_support"], weights),
            "feasibility": weighted_mean(g["feasibility"], weights),
        })
    return pd.DataFrame(rows)


def save_outputs(module_long, sub_long, anno_used, missing_df, constraint_df, outdir):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    module_metrics = [
        "activation", "direction", "confidence", "signed_flux", "ratio_support",
        "enzyme_capacity", "compartment_support", "cofactor_support", "feasibility",
        "mean_s3", "delta_s3"
    ]
    for metric in module_metrics:
        pivot_metric(module_long, metric).to_csv(outdir / f"module_{metric}_matrix.csv.gz", compression="gzip")

    sub_metrics = [
        "activation", "direction", "confidence", "signed_flux", "ratio_support",
        "enzyme_capacity", "compartment_support", "cofactor_support", "feasibility",
        "mean_s3", "delta_s3"
    ]
    for metric in sub_metrics:
        sub_long.pivot(index="subpathway_key", columns="celltype", values=metric).to_csv(
            outdir / f"subpathway_{metric}_matrix.csv.gz", compression="gzip"
        )

    module_long.to_csv(outdir / "module_scores_long.csv.gz", index=False, compression="gzip")
    sub_long.to_csv(outdir / "subpathway_scores_long.csv.gz", index=False, compression="gzip")
    anno_used.to_csv(outdir / "module_annotation_used.csv", index=False)
    missing_df.to_csv(outdir / "module_missing_trait_report.csv", index=False)
    constraint_df.to_csv(outdir / "module_constraint_gene_report.csv", index=False)

    summary = pd.DataFrame([{
        "annotation_rows": len(anno_used),
        "celltypes": int(module_long["celltype"].nunique()),
        "modules": int(module_long["module_id"].nunique()),
        "modules_with_ratio_assigned": int(anno_used["ratio_trait"].notna().sum()),
        "modules_with_enzyme_genes": int(anno_used["enzyme_genes"].apply(lambda x: isinstance(x, list) and len(x) > 0).sum()),
        "modules_missing_any_trait": int(missing_df["module_id"].nunique()) if not missing_df.empty else 0,
        "subpathways": int(sub_long["subpathway_key"].nunique()),
    }])
    summary.to_csv(outdir / "run_summary.csv", index=False)


def main():
    ap = argparse.ArgumentParser(description="Enzyme-constrained wAUCell-S3 + ratio metabolic flow model (round13)")
    ap.add_argument("--annotation", required=True, help="round12 annotation csv/xlsx")
    ap.add_argument("--score-matrix", required=True, help="wAUCell-S3 metabolite score matrix")
    ap.add_argument("--gene-matrix", required=True, help="gene expression matrix")
    ap.add_argument("--outdir", required=True)

    ap.add_argument("--enzyme-annotation", default=None,
                    help="optional round7 workbook or gene annotation table with enzyme genes")

    ap.add_argument("--metadata", default=None, help="optional cell metadata csv for aggregation")
    ap.add_argument("--id-col", default=None, help="row id column in metadata")
    ap.add_argument("--celltype-col", default=None, help="celltype column in metadata")
    ap.add_argument("--agg", default="mean", choices=["mean", "median"])

    ap.add_argument("--ratio-weight", type=float, default=1.0)
    ap.add_argument("--activation-scale", type=float, default=1.0)
    ap.add_argument("--direction-scale", type=float, default=1.0)
    ap.add_argument("--signal-scale", type=float, default=1.0)
    ap.add_argument("--ratio-scale", type=float, default=1.0)
    ap.add_argument("--enzyme-scale", type=float, default=1.0)
    ap.add_argument("--compartment-scale", type=float, default=1.0)
    ap.add_argument("--cofactor-scale", type=float, default=1.0)

    ap.add_argument("--ratio-missing-penalty", type=float, default=0.5)
    ap.add_argument("--no-ratio-factor", type=float, default=1.0)
    ap.add_argument("--no-enzyme-factor", type=float, default=0.8)
    ap.add_argument("--no-compartment-factor", type=float, default=0.9)
    ap.add_argument("--no-cofactor-factor", type=float, default=0.9)

    ap.add_argument("--enzyme-weight", type=float, default=1.0)
    ap.add_argument("--compartment-weight", type=float, default=0.5)
    ap.add_argument("--cofactor-weight", type=float, default=0.25)

    ap.add_argument("--gpr-mode", default="topk_mean", choices=["topk_mean", "mean", "max", "min", "geom"])
    ap.add_argument("--gpr-topk", type=int, default=3)

    args = ap.parse_args()

    anno = prepare_annotation(read_table(args.annotation))
    enz_map = build_enzyme_map(args.enzyme_annotation)
    anno = merge_annotation_with_enzymes(anno, enz_map)

    score_mat = read_matrix(args.score_matrix)
    gene_mat = read_matrix(args.gene_matrix)

    score_mat = aggregate_to_celltype(
        score_mat, metadata_path=args.metadata, id_col=args.id_col, celltype_col=args.celltype_col, agg=args.agg
    )
    gene_mat = aggregate_to_celltype(
        gene_mat, metadata_path=args.metadata, id_col=args.id_col, celltype_col=args.celltype_col, agg=args.agg
    )

    score_mat = score_mat.apply(pd.to_numeric, errors="coerce")
    gene_mat = gene_mat.apply(pd.to_numeric, errors="coerce")
    score_mat.columns = score_mat.columns.astype(str)
    score_mat.index = score_mat.index.astype(str)
    gene_mat.columns = gene_mat.columns.astype(str)
    gene_mat.index = gene_mat.index.astype(str)

    module_long, missing_df, constraint_df = compute_module_scores(
        anno=anno,
        score_mat=score_mat,
        gene_mat=gene_mat,
        ratio_weight=args.ratio_weight,
        activation_scale=args.activation_scale,
        direction_scale=args.direction_scale,
        signal_scale=args.signal_scale,
        ratio_scale=args.ratio_scale,
        enzyme_scale=args.enzyme_scale,
        compartment_scale=args.compartment_scale,
        cofactor_scale=args.cofactor_scale,
        ratio_missing_penalty=args.ratio_missing_penalty,
        no_ratio_factor=args.no_ratio_factor,
        no_enzyme_factor=args.no_enzyme_factor,
        no_compartment_factor=args.no_compartment_factor,
        no_cofactor_factor=args.no_cofactor_factor,
        enzyme_weight=args.enzyme_weight,
        compartment_weight=args.compartment_weight,
        cofactor_weight=args.cofactor_weight,
        gpr_mode=args.gpr_mode,
        gpr_topk=args.gpr_topk,
    )
    sub_long = aggregate_subpathway(module_long)
    save_outputs(module_long, sub_long, anno, missing_df, constraint_df, args.outdir)
    print("done")


if __name__ == "__main__":
    main()
